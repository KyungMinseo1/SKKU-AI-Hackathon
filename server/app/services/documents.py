import re
from dataclasses import dataclass
from pathlib import Path

MAX_SYLLABUS_CHARS = 30_000
MIN_PAGE_CHARS = 20
SPLIT_THRESHOLD = 1_500
TARGET_PIECE = 1_000


@dataclass
class ChunkDraft:
    page_from: int
    page_to: int
    text: str


def extract_pdf_markdown(path: str | Path) -> str:
    import pymupdf4llm

    return pymupdf4llm.to_markdown(str(path))[:MAX_SYLLABUS_CHARS]


def _pdf_pages(path: str) -> list[str]:
    import pymupdf4llm

    try:
        chunks = pymupdf4llm.to_markdown(path, page_chunks=True)
        if isinstance(chunks, list) and all(isinstance(c, dict) and "text" in c for c in chunks):
            return [str(c["text"]) for c in chunks]
    except Exception:  # noqa: BLE001 - fall back to plain pymupdf text
        pass
    import pymupdf

    with pymupdf.open(path) as doc:
        return [page.get_text("text", sort=True) for page in doc]


def _pptx_pages(path: str) -> list[str]:
    from pptx import Presentation

    pages: list[str] = []
    for slide in Presentation(path).slides:
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                parts.append(notes.text.strip())
        pages.append("\n".join(parts))
    return pages


def _split(text: str) -> list[str]:
    if len(text) <= SPLIT_THRESHOLD:
        return [text]
    pieces: list[str] = []
    current = ""
    for para in re.split(r"\n\s*\n", text):
        para = para.strip()
        if not para:
            continue
        if current and len(current) + len(para) + 2 > TARGET_PIECE:
            pieces.append(current)
            current = ""
        while len(para) > SPLIT_THRESHOLD:
            pieces.append(para[:TARGET_PIECE])
            para = para[TARGET_PIECE:]
        current = f"{current}\n\n{para}" if current else para
    if current:
        pieces.append(current)
    return pieces


def extract_chunks(path: str | Path, kind: str) -> list[ChunkDraft]:
    pages = _pdf_pages(str(path)) if kind == "pdf" else _pptx_pages(str(path))
    drafts: list[ChunkDraft] = []
    for page_no, text in enumerate(pages, 1):
        text = text.strip()
        if len(text) < MIN_PAGE_CHARS:
            continue
        for piece in _split(text):
            drafts.append(ChunkDraft(page_from=page_no, page_to=page_no, text=piece))
    return drafts
